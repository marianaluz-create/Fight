# -*- coding: utf-8 -*-
"""
ESQUELETO de build — copie este arquivo para cada novo cliente e preencha as
seções marcadas com "PREENCHER". Não roda "pronto" para um cliente novo: a
extração de cada aba (quais células, quais bullets grifar) é sempre um passo
manual/assistido — o que este script padroniza é a ESTRUTURA (as 15 seções,
o design system, o simulador, senha, autosave/export), não os dados.

Veja PLAYBOOK.md para o passo a passo completo.
"""
import openpyxl
from common import (
    fmt, fmt_shortdate, fmt_fulldate, fmt_monyear, esc, escp, commafy,
    ul_cur, ul_para, ul_lines, table, table_h, sheet_rows, colored_sheets,
    PaneBuilder,
)

# ========== 1) CONFIGURAÇÃO DO CLIENTE (PREENCHER) ==========
CLIENT_NAME = "Nome do Cliente"
YEAR = "2026"
SENHA = f"{CLIENT_NAME}{YEAR}"                 # regra fixa: Nome do Cliente + Ano Corrente
DATA_APRESENTACAO = "10 de julho de 2026"
XLSX_PATH = "/caminho/para/planilha_do_cliente.xlsx"

wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

# ========== 2) CONFIRMAR AS ABAS-FONTE PELA COR ==========
# Regra do cliente (O2): as abas do relatório são sempre as que estão
# pintadas na planilha (tabColor) — normalmente com o verde de marca da O2.
# Rode isto primeiro e confira a lista antes de mapear qualquer célula.
print("Abas coloridas encontradas (candidatas a virar seção do HTML):")
for name, rgb in colored_sheets(wb):
    print(" -", name, rgb)
print("\nAbas SEM cor (normalmente apoio/rascunho — Base, Apoio, Histórico...):")
for name in wb.sheetnames:
    if name not in [n for n, _ in colored_sheets(wb)]:
        print(" -", name)
# Atenção: confira manualmente se alguma aba usada no relatório ficou sem
# cor por esquecimento (já aconteceu com "Projeção 13 Semanas" e "Relação de
# Endividamento" no arquivo da Fight Analytics) — se sim, use-a mesmo assim,
# mas avise o cliente para colorir da próxima vez.

pb = PaneBuilder()

# ========== 3) UMA SEÇÃO POR ABA (PREENCHER — repita o padrão abaixo) ==========
# Tabs 1-6 (Capa, Sumário, Disclaimer, Fonte de Dados, Cenário Macro, Panorama)
# são majoritariamente narrativas/institucionais — normalmente não vêm de uma
# tabela de planilha, e sim de texto livre + pesquisa externa (site do
# cliente, setor). Escreva-as à mão seguindo o padrão visual das abas 7+.

# Exemplo real (aba 7 — DRE Gerencial), adaptar células/rótulos para o cliente:
#
# s = wb['7.1 DRE Gerencial']
# dre_h = [s[f"{c}7"].value for c in ['B','C','D','E','F','G','H','I','J']]
# dre_rows = sheet_rows(wb, '7.1 DRE Gerencial', 8, 35, ['B','C','D','E','F','G','H','I','J'])
# body = f'''      <h2 class="section-title">Resultado</h2>
#       <h3 class="section-h">DRE Gerencial</h3>
#       <p class="lead" style="margin-bottom:2px">{escp(s["B5"].value)}</p>
#       <div class="split tmaj">
#         <div class="card">{table(dre_h, dre_rows, green_key=True)}</div>
#       </div>'''
# pb.add('dre', '7', 'DRE Gerencial', body)
#
# ... repita para cada uma das 15 abas, na ordem definida no PLAYBOOK.md ...

# ========== 4) SIMULADOR (aba Projeção de Caixa, 13 Semanas) ==========
# O corpo HTML da aba (inputs, tabelas) segue o mesmo padrão de table()/PaneBuilder
# acima. Os NÚMEROS HISTÓRICOS do motor JS do simulador (CUM_RECV, CUM_PAY, D21,
# F21) NÃO vêm daqui — são calculados a partir da aba transacional "Base" do
# cliente (ver PLAYBOOK.md, seção "Simulador — como recalcular a base
# histórica") e injetados nos placeholders {{SIM_CUM_RECV}} / {{SIM_CUM_PAY}} /
# {{SIM_D21}} / {{SIM_F21}} do shell.html, abaixo.
SIM_CUM_RECV = "[]"   # PREENCHER — ver PLAYBOOK.md
SIM_CUM_PAY = "[]"    # PREENCHER — ver PLAYBOOK.md
SIM_D21 = "0"         # PREENCHER — referência histórica, método Pior Mês
SIM_F21 = "0"         # PREENCHER — referência histórica, método Dias de Pagamento

# ========== 5) MONTAR O ARQUIVO FINAL A PARTIR DO shell.html ==========
shell = open('shell.html', encoding='utf-8').read()
final = (shell
    .replace('{{CLIENT_NAME}}', CLIENT_NAME)
    .replace('{{DATA_APRESENTACAO}}', DATA_APRESENTACAO)
    .replace('{{SENHA}}', SENHA)
    .replace('{{NAV_BUTTONS}}', pb.render_navs())
    .replace('{{NAV_PANES}}', pb.render_panes())
    .replace('{{SIM_CUM_RECV}}', SIM_CUM_RECV)
    .replace('{{SIM_CUM_PAY}}', SIM_CUM_PAY)
    .replace('{{SIM_D21}}', SIM_D21)
    .replace('{{SIM_F21}}', SIM_F21)
)
outname = f"diagnostico-{CLIENT_NAME.lower().replace(' ', '-')}.html"
open(outname, 'w', encoding='utf-8').write(final)
print("\nOK ->", outname, len(final), "bytes")
